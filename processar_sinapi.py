#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Processador de Tabelas Oficiais SINAPI (Caixa Econômica Federal)
Converte planilhas brutas (.xlsx) em catálogo leve e otimizado para o GeoGestor.
"""

import os
import re
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DADOS_DIR = os.path.join(BASE_DIR, "dados_orcamento")
OUTPUT_DIR = os.path.join(BASE_DIR, "data", "sinapi")

def extrair_codigo_hyperlink(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return int(val)
    val_str = str(val).strip()
    # Buscar sequencias de digitos
    m = re.findall(r'\d+', val_str)
    if m:
        for cand in reversed(m):
            if len(cand) >= 4:
                return int(cand)
        return int(m[-1])
    return None

def processar_sinapi(uf="PB", mes_ref="2026_07"):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    arquivo_ref = os.path.join(DADOS_DIR, f"SINAPI_Referência_{mes_ref}.xlsx")
    
    if not os.path.exists(arquivo_ref):
        print(f"Arquivo não encontrado: {arquivo_ref}")
        return False
        
    print(f"Carregando {arquivo_ref} (read_only=True)...")
    wb = openpyxl.load_workbook(arquivo_ref, read_only=True, data_only=False)
    
    # 1. PROCESSAR CCD (Composições Com Desoneração) e CSD (Sem Desoneração)
    print("Processando CCD e CSD...")
    ws_ccd = wb['CCD']
    ws_csd = wb['CSD']
    
    col_pb_ccd = 32
    col_pb_csd = 32
    
    r_count = 0
    composicoes = {}
    
    for row in ws_ccd.iter_rows(values_only=True):
        r_count += 1
        if r_count == 9:
            for idx, val in enumerate(row):
                if val == uf:
                    col_pb_ccd = idx
            print(f"Coluna {uf} no CCD: {col_pb_ccd}")
        elif r_count >= 11:
            grupo = row[0]
            cod_raw = row[1]
            desc = row[2]
            und = row[3]
            custo_des = row[col_pb_ccd] if col_pb_ccd < len(row) else None
            
            codigo = extrair_codigo_hyperlink(cod_raw)
            if codigo and desc:
                try:
                    c_des = float(custo_des) if custo_des is not None and str(custo_des).strip() not in ('', '-') else 0.0
                except:
                    c_des = 0.0
                composicoes[codigo] = {
                    "codigo": str(codigo),
                    "tipo": "COMPOSICAO",
                    "grupo": str(grupo).strip() if grupo else "",
                    "descricao": str(desc).strip(),
                    "unidade": str(und).strip() if und else "UN",
                    "preco_desonerado": round(c_des, 2),
                    "preco_nao_desonerado": round(c_des, 2)
                }
    
    # Atualizar preço sem desoneração do CSD
    r_count = 0
    for row in ws_csd.iter_rows(values_only=True):
        r_count += 1
        if r_count == 9:
            for idx, val in enumerate(row):
                if val == uf:
                    col_pb_csd = idx
            print(f"Coluna {uf} no CSD: {col_pb_csd}")
        elif r_count >= 11:
            cod_raw = row[1]
            custo_sem = row[col_pb_csd] if col_pb_csd < len(row) else None
            codigo = extrair_codigo_hyperlink(cod_raw)
            if codigo and codigo in composicoes:
                try:
                    c_sem = float(custo_sem) if custo_sem is not None and str(custo_sem).strip() not in ('', '-') else composicoes[codigo]["preco_desonerado"]
                except:
                    c_sem = composicoes[codigo]["preco_desonerado"]
                composicoes[codigo]["preco_nao_desonerado"] = round(c_sem, 2)
                
    print(f"Total de composições processadas para {uf}: {len(composicoes)}")
    
    # 2. PROCESSAR INSUMOS (ICD e ISD)
    print("Processando ICD e ISD...")
    ws_icd = wb['ICD']
    ws_isd = wb['ISD']
    
    col_pb_icd = 19
    col_pb_isd = 19
    insumos = {}
    
    r_count = 0
    for row in ws_icd.iter_rows(values_only=True):
        r_count += 1
        if r_count == 10:
            for idx, val in enumerate(row):
                if val == uf:
                    col_pb_icd = idx
            print(f"Coluna {uf} no ICD: {col_pb_icd}")
        elif r_count >= 11:
            classif = row[0]
            cod_raw = row[1]
            desc = row[2]
            und = row[3]
            preco_des = row[col_pb_icd] if col_pb_icd < len(row) else None
            
            codigo = extrair_codigo_hyperlink(cod_raw)
            if codigo and desc:
                try:
                    p_des = float(preco_des) if preco_des is not None and str(preco_des).strip() not in ('', '-') else 0.0
                except:
                    p_des = 0.0
                insumos[codigo] = {
                    "codigo": str(codigo),
                    "tipo": "INSUMO",
                    "grupo": str(classif).strip() if classif else "MATERIAL",
                    "descricao": str(desc).strip(),
                    "unidade": str(und).strip() if und else "UN",
                    "preco_desonerado": round(p_des, 2),
                    "preco_nao_desonerado": round(p_des, 2)
                }
                
    r_count = 0
    for row in ws_isd.iter_rows(values_only=True):
        r_count += 1
        if r_count == 10:
            for idx, val in enumerate(row):
                if val == uf:
                    col_pb_isd = idx
        elif r_count >= 11:
            cod_raw = row[1]
            preco_sem = row[col_pb_isd] if col_pb_isd < len(row) else None
            codigo = extrair_codigo_hyperlink(cod_raw)
            if codigo and codigo in insumos:
                try:
                    p_sem = float(preco_sem) if preco_sem is not None and str(preco_sem).strip() not in ('', '-') else insumos[codigo]["preco_desonerado"]
                except:
                    p_sem = insumos[codigo]["preco_desonerado"]
                insumos[codigo]["preco_nao_desonerado"] = round(p_sem, 2)
                
    print(f"Total de insumos processados para {uf}: {len(insumos)}")
    
    # 3. COMPOR CATÁLOGO UNIFICADO PARA BUSCA RÁPIDA (AutoComplete)
    catalogo = []
    for c in composicoes.values():
        catalogo.append(c)
    for i in insumos.values():
        catalogo.append(i)
        
    print(f"Total geral no catálogo {uf}: {len(catalogo)} itens")
    
    # 4. PROCESSAR ANALÍTICO (Relação Pai -> Filhos)
    print("Processando aba Analítico...")
    ws_analitico = wb['Analítico']
    analitico = {}
    r_count = 0
    for row in ws_analitico.iter_rows(values_only=True):
        r_count += 1
        if r_count >= 11:
            cod_pai_raw = row[1]
            tipo_item = row[2]
            cod_filho_raw = row[3]
            desc_filho = row[4]
            und_filho = row[5]
            coef_raw = row[6]
            
            cod_pai = extrair_codigo_hyperlink(cod_pai_raw)
            if not cod_pai:
                continue
                
            cod_pai_str = str(cod_pai)
            if cod_pai_str not in analitico:
                analitico[cod_pai_str] = []
                
            if tipo_item and desc_filho:
                try:
                    coef = float(coef_raw) if coef_raw is not None else 1.0
                except:
                    coef = 1.0
                analitico[cod_pai_str].append({
                    "tipo": str(tipo_item).strip(),
                    "codigo": str(extrair_codigo_hyperlink(cod_filho_raw) or ""),
                    "descricao": str(desc_filho).strip(),
                    "unidade": str(und_filho).strip() if und_filho else "",
                    "coeficiente": coef
                })
                
    print(f"Composições detalhadas no Analítico: {len(analitico)}")
    
    # SALVAR ARQUIVOS JSON OTIMIZADOS
    arquivo_cat = os.path.join(OUTPUT_DIR, f"catalogo_{uf.lower()}_{mes_ref}.json")
    with open(arquivo_cat, "w", encoding="utf-8") as f:
        json.dump(catalogo, f, ensure_ascii=False, indent=None)
        
    arquivo_ana = os.path.join(OUTPUT_DIR, f"analitico_{mes_ref}.json")
    with open(arquivo_ana, "w", encoding="utf-8") as f:
        json.dump(analitico, f, ensure_ascii=False, indent=None)
        
    metadados = {
        "uf": uf,
        "mes_referencia": "07/2026",
        "chave_referencia": f"{uf.upper()}_2026_07",
        "nome_exibicao": f"SINAPI 07/2026 - {uf.upper()} (João Pessoa)",
        "data_emissao": "11/08/2026",
        "encargos_sociais": {
            "horista": 1.0206,
            "mensalista": 0.5991
        },
        "total_composicoes": len(composicoes),
        "total_insumos": len(insumos),
        "total_itens": len(catalogo),
        "arquivo_catalogo": f"catalogo_{uf.lower()}_{mes_ref}.json",
        "arquivo_analitico": f"analitico_{mes_ref}.json"
    }
    
    arquivo_meta = os.path.join(OUTPUT_DIR, f"metadados_{uf.lower()}_{mes_ref}.json")
    with open(arquivo_meta, "w", encoding="utf-8") as f:
        json.dump(metadados, f, ensure_ascii=False, indent=2)
        
    # Salvar índice mestre de bases disponíveis
    bases_index_file = os.path.join(OUTPUT_DIR, "bases_disponiveis.json")
    bases = []
    if os.path.exists(bases_index_file):
        try:
            with open(bases_index_file, "r", encoding="utf-8") as f:
                bases = json.load(f)
        except:
            bases = []
            
    # Atualizar ou inserir
    bases = [b for b in bases if b.get("chave_referencia") != metadados["chave_referencia"]]
    bases.append(metadados)
    with open(bases_index_file, "w", encoding="utf-8") as f:
        json.dump(bases, f, ensure_ascii=False, indent=2)
        
    print(f"\n[SUCESSO] Processamento concluído com êxito!")
    print(f"Catálogo: {arquivo_cat} ({os.path.getsize(arquivo_cat) / (1024*1024):.2f} MB)")
    print(f"Analítico: {arquivo_ana} ({os.path.getsize(arquivo_ana) / (1024*1024):.2f} MB)")
    print(f"Metadados: {arquivo_meta}")
    return True

if __name__ == "__main__":
    processar_sinapi(uf="PB", mes_ref="2026_07")
