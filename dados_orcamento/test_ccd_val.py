import openpyxl

path = r"c:\Users\Windows 11\Documents\Projetos\Projeto_V01\Projetos_Obras\dados_orcamento\SINAPI_Referência_2026_07.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
ws = wb['CCD']
r = list(ws.iter_rows(values_only=True))[10]
print("data_only=False, col 32:", r[32], type(r[32]))

wb_data = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws_data = wb_data['CCD']
r_data = list(ws_data.iter_rows(values_only=True))[10]
print("data_only=True, col 1:", r_data[1], type(r_data[1]))
print("data_only=True, col 32:", r_data[32], type(r_data[32]))
