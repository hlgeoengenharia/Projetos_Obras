import os
import openpyxl
import xlrd

DIR = r"c:\Users\Windows 11\Documents\Projetos\Projeto_V01\Projetos_Obras\dados_orcamento"

def inspect_xlsx(filename):
    path = os.path.join(DIR, filename)
    print(f"\n{'='*60}\nFILE: {filename}\n{'='*60}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)
    for sheet_name in wb.sheetnames[:3]: # inspect first sheets
        print(f"\n--- Sheet: {sheet_name} ---")
        ws = wb[sheet_name]
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if any(row):
                row_count += 1
                if row_count <= 8:
                    print(f"Row {row_count}: {row[:12]}")
                elif row_count == 9:
                    print("...")
        print(f"Total non-empty sample rows seen: {row_count}")

def inspect_xls(filename):
    path = os.path.join(DIR, filename)
    print(f"\n{'='*60}\nFILE: {filename}\n{'='*60}")
    try:
        wb = xlrd.open_workbook(path)
        print("Sheets:", wb.sheet_names())
        for name in wb.sheet_names()[:5]:
            print(f"\n--- Sheet: {name} ---")
            sh = wb.sheet_by_name(name)
            print(f"Rows: {sh.nrows}, Cols: {sh.ncols}")
            for r in range(min(10, sh.nrows)):
                row_vals = sh.row_values(r)
                if any(row_vals):
                    print(f"Row {r}: {row_vals[:10]}")
    except Exception as e:
        print(f"Error reading {filename} with xlrd: {e}")
        # Try openpyxl if it was actually xlsx renamed to xls
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            print("Opened with openpyxl! Sheets:", wb.sheetnames)
        except Exception as e2:
            print(f"Also failed with openpyxl: {e2}")

if __name__ == "__main__":
    for f in ["SINAPI_Referência_2026_07.xlsx", "SINAPI_Manutenções_2026_07.xlsx", "SINAPI_familias_e_coeficientes_2026_07.xlsx", "SINAPI_mao_de_obra_2026_07.xlsx"]:
        inspect_xlsx(f)
    inspect_xls("BASE MODELO ORÇAMENTO.xls")
