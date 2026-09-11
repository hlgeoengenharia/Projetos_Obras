import xlrd
import openpyxl
import os

DIR = r"c:\Users\Windows 11\Documents\Projetos\Projeto_V01\Projetos_Obras\dados_orcamento"

def analyze_modelo():
    path = os.path.join(DIR, "BASE MODELO ORÇAMENTO.xls")
    wb = xlrd.open_workbook(path, formatting_info=False)
    print("=== MODELO MUNICÍPIO: BASE MODELO ORÇAMENTO.xls ===")
    print("Abas:", wb.sheet_names())
    
    for sname in ["QUADRO_RESUMO", "MC_ANEXO SEMOB", "ORÇ_ANEXO SEMOB_DES", "CPUs_DES", "Cronograma", "CURVA_ABC", "BDI_DES_EDIFICAÇÃO"]:
        if sname in wb.sheet_names():
            sh = wb.sheet_by_name(sname)
            print(f"\n--- ABA: {sname} (Linhas: {sh.nrows}, Colunas: {sh.ncols}) ---")
            for r in range(min(25, sh.nrows)):
                row = [str(x).strip() for x in sh.row_values(r) if str(x).strip() != ""]
                if row:
                    print(f"L{r+1}: {' | '.join(row[:8])}")

def analyze_sinapi_ref():
    path = os.path.join(DIR, "SINAPI_Referência_2026_07.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("\n=== SINAPI REFERÊNCIA 2026_07 ===")
    for sname in ['ICD', 'CCD', 'Analítico', 'Analítico com Custo']:
        if sname in wb.sheetnames:
            ws = wb[sname]
            print(f"\n--- ABA: {sname} ---")
            count = 0
            for row in ws.iter_rows(values_only=True):
                if any(row):
                    count += 1
                    if count <= 7:
                        clean_row = [str(x) if x is not None else "" for x in row[:10]]
                        print(f"L{count}: {' | '.join(clean_row)}")
                    elif count == 8:
                        print("...")
            print(f"Total de linhas preenchidas na aba {sname}: {count}")

if __name__ == "__main__":
    analyze_modelo()
    analyze_sinapi_ref()
